# 生成一份真实结构的 pptx，用来验证插件的手写解析器。
#
# 为什么不把生成结果提交进仓库：真实 PPT 里常含客户名称、报价、项目细节。
# 这里只造一份结构真实（含母版、备注、多段落、空页）但内容无关痛痒的文件，
# 生成到临时目录，验完即删。
#
# 运行：python tools/make-fixture-pptx.py <输出路径>

import sys
from pptx import Presentation
from pptx.util import Inches, Pt

out = sys.argv[1] if len(sys.argv) > 1 else 'fixture.pptx'

prs = Presentation()

# 第 1 页：标题 + 副标题（真实 PPT 常见结构）
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = '半自动打码贴标线技术方案'
slide.placeholders[1].text = 'Data Matrix 追溯 · 产线集成'

# 第 2 页：多级项目符号（验证段落还原）
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = '设备构成'
body = slide.placeholders[1].text_frame
body.text = '输送线体'
for line, level in [('入口段 / 工艺段 / 出口段', 1), ('导向条与止动器', 1)]:
    p = body.add_paragraph()
    p.text = line
    p.level = level
for line in ['打印贴标机头（热转印 300dpi）', '视觉检测：2 套相机 + 光源', '剔除机构与不良品箱']:
    p = body.add_paragraph()
    p.text = line
    p.level = 0

# 第 3 页：表格（验证表格文字也能被读出）
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = '主要部件清单'
rows, cols = 3, 2
table = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.8), Inches(8), Inches(1.8)).table
data = [['部件', '说明'], ['打印贴标机', '热转印 300dpi'], ['工业相机', '读码 + 复检']]
for r in range(rows):
    for c in range(cols):
        table.cell(r, c).text = data[r][c]

# 第 4 页：横排文本框（真实 PPT 常有多个文本框，且不在占位符里）
slide = prs.slides.add_slide(prs.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(7), Inches(1.2))
box.text_frame.text = '换行测试第一行'
box.text_frame.add_paragraph().text = '换行测试第二行'

# 第 5 页：只有图片、没有文字（验证"没有可提取文字"的提示）
slide = prs.slides.add_slide(prs.slide_layouts[6])
png_path = out + '.probe.png'
try:
    import struct
    import zlib

    def chunk(tag, data):
        c = tag + data
        return struct.pack('>I', len(data)) + c + struct.pack('>I', zlib.crc32(c) & 0xffffffff)

    raw = b''.join(b'\x00' + b'\xff\xff\xff' * 8 for _ in range(8))
    png = (b'\x89PNG\r\n\x1a\n'
           + chunk(b'IHDR', struct.pack('>IIBBBBB', 8, 8, 8, 2, 0, 0, 0))
           + chunk(b'IDAT', zlib.compress(raw))
           + chunk(b'IEND', b''))
    with open(png_path, 'wb') as f:
        f.write(png)
    slide.shapes.add_picture(png_path, Inches(1), Inches(1), Inches(2), Inches(2))
except Exception as e:  # 图片造不出来也不影响其它页的验证
    print(f'  （跳过图片页：{e}）', file=sys.stderr)

# 备注：第 2 页写一条有信息量的，第 3 页写一条纯页码占位（应被丢弃）
prs.slides[1].notes_slide.notes_text_frame.text = '本页重点是产能与节拍，客户最关心'
prs.slides[2].notes_slide.notes_text_frame.text = '3'

prs.save(out)
print(f'已生成：{out}')
print(f'页数：{len(prs.slides.__iter__.__self__._sldIdLst)}')
